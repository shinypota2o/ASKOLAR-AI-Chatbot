"""Export the 'Training Questions' sheet of the ASKOLAR Excel file to training/intent_dataset.csv.

Usage:
    python training/export_from_excel.py "path/to/ASKOLAR.xlsx"
"""
import csv
import sys
from pathlib import Path

from openpyxl import load_workbook

COLUMNS = ['question_id', 'text', 'label', 'language', 'contributor', 'review_status', 'paraphrase_group']
HEADER_MAP = {
    'question id': 'question_id', 'question': 'text', 'intent': 'label', 'language': 'language',
    'contributor': 'contributor', 'review status': 'review_status', 'paraphrase group': 'paraphrase_group',
}


def main(xlsx_path, out_path=Path(__file__).with_name('intent_dataset.csv')):
    sheet = load_workbook(xlsx_path, data_only=True)['Training Questions']
    rows = list(sheet.iter_rows(values_only=True))
    header = [HEADER_MAP.get(str(h).strip().lower()) if h else None for h in rows[0]]
    written = 0
    with open(out_path, 'w', newline='', encoding='utf-8') as f:
        writer = csv.DictWriter(f, fieldnames=COLUMNS)
        writer.writeheader()
        for number, row in enumerate(rows[1:], start=1):
            record = {key: (str(value).strip() if value is not None else '') for key, value in zip(header, row) if key}
            if not record.get('text') or not record.get('label'):
                continue
            # Question IDs come from an Excel formula; rebuild them if the cached value is missing.
            record['question_id'] = record.get('question_id') or f'Q{number:04d}'
            writer.writerow({c: record.get(c, '') for c in COLUMNS})
            written += 1
    print(f'Wrote {written} questions to {out_path}')


if __name__ == '__main__':
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
